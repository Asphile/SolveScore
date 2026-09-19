"""Best-effort text extraction from uploaded documents, for AI-content-detection
purposes. Video/audio transcription is not implemented in this prototype -- see
extract_text_from_video for the documented integration point.
"""

import os


def extract_text(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext == ".pdf":
            return _extract_pdf(file_path)
        if ext == ".docx":
            return _extract_docx(file_path)
        if ext == ".pptx":
            return _extract_pptx(file_path)
    except Exception:
        return ""
    return ""


def _extract_pdf(file_path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(file_path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(file_path: str) -> str:
    from docx import Document

    document = Document(file_path)
    return "\n".join(p.text for p in document.paragraphs)


def _extract_pptx(file_path: str) -> str:
    from pptx import Presentation

    presentation = Presentation(file_path)
    lines = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.append(shape.text_frame.text)
    return "\n".join(lines)


def extract_text_from_video(file_path: str) -> str:
    """Integration point for video transcription (speech-to-text).

    Not implemented in this prototype -- wire up a transcription provider
    (configured via environment variables, like the AI detection provider)
    before relying on video-transcript AI-content analysis.
    """
    raise NotImplementedError("Video transcription is not configured. See app/services/text_extraction.py.")
